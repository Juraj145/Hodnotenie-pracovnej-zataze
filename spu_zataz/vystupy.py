"""Export výsledkov do Excelu."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calc import Obdobie, VysledokUcitela, VysledokUstavov
from .version import __version__

STLPCE_UCITELIA = [
    ("Meno", lambda r: r.ucitel.cele_meno, 30),
    ("Osobné číslo", lambda r: r.ucitel.osobne_cislo, 12),
    ("Fakulta", lambda r: r.ucitel.fakulta, 10),
    ("Ústav", lambda r: r.ucitel.ustav, 22),
    ("Funkcia", lambda r: r.ucitel.funkcia, 16),
    ("Úväzok", lambda r: r.ucitel.uvazok, 8),
    ("Fond (h/rok)", lambda r: r.fond_hodin, 10),
    ("Výučba (h)", lambda r: r.h_vyucba, 10),
    ("Študenti (h)", lambda r: r.h_studenti, 10),
    ("Záverečné práce (h)", lambda r: r.h_zaverecne_prace, 11),
    ("Vzdelávanie spolu (h/ak. rok)", lambda r: r.h_vzdelavanie, 13),
    ("Vzdelávanie % fondu", lambda r: r.pct_vzdelavanie / 100, 11),
    ("Projekty (h/rok)", lambda r: r.h_projekty, 10),
    ("Spolu % fondu", lambda r: r.pct_spolu / 100, 10),
    ("Status (čl. 7)", lambda r: r.status, 34),
    ("Priama výučba h/týždeň", lambda r: r.vyucba_tyzden, 11),
    ("Referencia h/týždeň", lambda r: r.referencna_vyucba_tyzden, 11),
    ("Publikácie – body/rok", lambda r: r.body_publikacie, 12),
    ("Projekty – podiel na financiách €/rok", lambda r: r.financie_projekty, 14),
    ("Štand. vzdelávanie", lambda r: r.std_vzdelavanie, 11),
    ("Štand. publikácie", lambda r: r.std_publikacie, 11),
    ("Štand. projekty", lambda r: r.std_projekty, 11),
    ("Upozornenia", lambda r: "; ".join(r.upozornenia), 50),
]

STLPCE_USTAVY = [
    ("Fakulta", lambda u: u.fakulta, 10),
    ("Ústav", lambda u: u.ustav, 28),
    ("Počet učiteľov", lambda u: u.pocet_ucitelov, 9),
    ("Prepočítané úväzky", lambda u: u.uvazky, 11),
    ("Výkon vzdelávanie (študentohodiny × koef.)", lambda u: u.vykon_vzdelavanie, 16),
    ("Výkon publikácie (body)", lambda u: u.vykon_publikacie, 13),
    ("Výkon projekty (€)", lambda u: u.vykon_projekty, 13),
    ("z vzdelávanie", lambda u: u.z_vzdelavanie, 10),
    ("z publikácie", lambda u: u.z_publikacie, 10),
    ("z projekty", lambda u: u.z_projekty, 10),
    ("Sumárne skóre (40/40/20)", lambda u: u.sumarne_skore, 12),
    ("Optimum vzdelávanie", lambda u: u.opt_vzdelavanie, 11),
    ("Optimum publikácie", lambda u: u.opt_publikacie, 11),
    ("Optimum projekty", lambda u: u.opt_projekty, 11),
    ("Kvázi optimálny počet pedagógov", lambda u: u.opt_pedagogovia, 13),
    ("Rozdiel oproti stavu", lambda u: u.rozdiel_pedagogov, 11),
]

FARBY_STATUSU = {
    "Ideálny": "C6EFCE",
    "Pod ideálnym": "DDEBF7",
    "Nad ideálnym": "FFF2CC",
    "Nutné": "FCE4D6",
    "Preťaženie": "FFC7CE",
    "Nezahrnutý": "EDEDED",
}

_HF = PatternFill("solid", fgColor="1E4620")
_HFONT = Font(bold=True, color="FFFFFF")


def _list(ws, stlpce, riadky, pct_cols=()):
    for j, (h, _, w) in enumerate(stlpce, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill, c.font = _HF, _HFONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 45
    for i, r in enumerate(riadky, start=2):
        for j, (h, f, _) in enumerate(stlpce, start=1):
            v = f(r)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                c.number_format = "0.0%" if h in pct_cols else "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def export_vysledkov(path: str | Path, ucitelia: list[VysledokUcitela], ustavy: VysledokUstavov,
                     obd: Obdobie, params: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Učitelia"
    _list(ws, STLPCE_UCITELIA, ucitelia, pct_cols=("Vzdelávanie % fondu", "Spolu % fondu"))
    status_col = [h for h, _, _ in STLPCE_UCITELIA].index("Status (čl. 7)") + 1
    for i in range(2, len(ucitelia) + 2):
        c = ws.cell(row=i, column=status_col)
        for k, farba in FARBY_STATUSU.items():
            if str(c.value).startswith(k):
                c.fill = PatternFill("solid", fgColor=farba)

    ws2 = wb.create_sheet("Ústavy")
    _list(ws2, STLPCE_USTAVY, ustavy.ustavy)

    ws3 = wb.create_sheet("Model a parametre")
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 90
    info = [
        ("Vygenerované", dt.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Verzia programu", __version__),
        ("Metodika", "Metodický pokyn 1/2023 v znení Dodatku č. 2 (účinný od 1. 7. 2025)"),
        ("Obdobie", obd.popis()),
    ]
    for oblast, reg in ustavy.regresie.items():
        info.append((f"Regresia {oblast}", f"y = {reg.b:.4f}·x ; sm. odchýlka rezíduí {reg.s:.4f}" if reg else "nedostatok údajov"))
    for u in ustavy.upozornenia:
        info.append(("Upozornenie", u))
    info.append(("Parametre (JSON)", json.dumps(params, ensure_ascii=False)))
    for k, v in info:
        ws3.append([k, v])
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
        ws3.cell(row=ws3.max_row, column=2).alignment = Alignment(wrap_text=True)
    wb.save(path)
