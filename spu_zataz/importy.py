"""Import a export údajov: Excel šablóna, exporty z UIS / CREPČ (CSV, XLSX) s mapovaním stĺpcov."""

from __future__ import annotations

import csv
import json
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config
from .db import Databaza
from .models import (Projekt, ProjektUcast, Publikacia, Ucitel, Vyucba,
                     ZaverecnaPraca)


# ------------------------------------------------------------------ pomocné prevody

def norm(text) -> str:
    """Normalizácia názvu stĺpca: malé písmená, bez diakritiky a nadbytočných medzier."""
    s = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("_", " ").split())


def to_float(v, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace("€", "").replace("%", "")
    if s.count(",") == 1 and s.count(".") >= 1:   # 1.234,56
        s = s.replace(".", "").replace(",", ".")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


def to_uvazok(v) -> float:
    x = to_float(v, 1.0)
    return x / 100 if x > 1.5 else x   # 100 → 1,0 ; 50 → 0,5


def to_int(v, default: int = 0) -> int:
    return int(round(to_float(v, default)))


def to_bool(v) -> bool:
    return norm(v) in {"ano", "a", "1", "true", "x", "yes", "y", "zodpovedny", "zodpovedny riesitel", "ved"}


def to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_ak_rok(v) -> str:
    s = to_str(v).replace(" ", "").replace("-", "/")
    if "/" in s:
        a, b = s.split("/", 1)
        if len(b) == 2 and len(a) == 4:
            b = a[:2] + b
        return f"{a}/{b}"
    if s.isdigit() and len(s) == 4:   # 2024 → 2024/2025
        return f"{s}/{int(s) + 1}"
    return s


def to_jazyk(v) -> str:
    s = norm(v)
    if s.startswith("mob") or "erasmus" in s:
        return "MOB"
    if s in {"en", "eng", "anglicky", "anglictina", "english", "aj"} or s.startswith("angl") or "english" in s:
        return "EN"
    return "SK"


def to_stupen(v) -> str:
    s = norm(v)
    if s.startswith("3") or "phd" in s or "dok" in s or "dizert" in s:
        return "PhD"
    if s.startswith("2") or "ing" in s or "mgr" in s or "dipl" in s:
        return "Ing"
    return "Bc"


def to_kvartil(v) -> str:
    s = norm(v).upper().replace(" ", "")
    for q in ("Q1", "Q2", "Q3", "Q4"):
        if q in s:
            return q
    if s in {"1", "2", "3", "4"}:
        return "Q" + s
    return "Bez Q"


def to_kategoria(v) -> str:
    s = norm(v)
    for k in config.KATEGORIE_PUBLIKACII:
        if norm(k) == s:
            return k
    if "monograf" in s:
        return config.KATEGORIE_PUBLIKACII[0]
    if "edit" in s or "kritick" in s:
        return config.KATEGORIE_PUBLIKACII[1]
    if ("wos" in s or "scopus" in s or "index" in s or "current contents" in s) and "ostat" not in s:
        return config.KATEGORIE_PUBLIKACII[2]
    return config.KATEGORIE_PUBLIKACII[3]


def to_podiel(v) -> float:
    x = to_float(v, 1.0)
    return x / 100 if x > 1.0 else x   # 25 % → 0,25


# ------------------------------------------------------------------ špecifikácia stĺpcov

@dataclass
class Stlpec:
    pole: str                  # interný názov
    hlavicka: str              # hlavička v šablóne
    prevod: Callable = to_str
    povinny: bool = False
    synonyma: tuple = ()       # alternatívne názvy stĺpcov (UIS, CREPČ)


# Polia identifikujúce učiteľa – spoločné pre všetky typy
_ID_UCITELA = [
    Stlpec("osobne_cislo", "Osobné číslo", to_str, False, ("osobne cislo", "os. cislo", "os cislo", "id osoby", "id zamestnanca", "uco", "osc")),
    Stlpec("meno", "Meno a priezvisko", to_str, False, ("meno", "priezvisko a meno", "ucitel", "vyucujuci", "autor", "zamestnanec", "riesitel", "veduci prace", "skolitel", "pedagog")),
]


@dataclass
class TypImportu:
    kluc: str
    nazov: str                 # názov hárku v šablóne
    stlpce: list[Stlpec]
    popis: str = ""


TYPY: dict[str, TypImportu] = {
    "ucitelia": TypImportu("ucitelia", "Učitelia", [
        _ID_UCITELA[0],
        Stlpec("titul_pred", "Titul pred menom", to_str, False, ("titul pred", "tituly pred")),
        Stlpec("meno", "Meno a priezvisko", to_str, True, ("meno", "priezvisko a meno", "zamestnanec")),
        Stlpec("titul_za", "Titul za menom", to_str, False, ("titul za", "tituly za")),
        Stlpec("fakulta", "Fakulta", to_str, False, ()),
        Stlpec("ustav", "Ústav", to_str, False, ("pracovisko", "katedra", "ustav/katedra")),
        Stlpec("funkcia", "Funkcia", to_str, False, ("funkcne miesto", "pozicia", "kategoria")),
        Stlpec("uvazok", "Úväzok (0–1)", to_uvazok, False, ("uvazok", "prepocitany uvazok", "uvazok %", "fte")),
        Stlpec("podiel_aktivny", "Podiel obdobia v PP (0–1)", lambda v: to_float(v, 1.0), False,
               ("podiel obdobia", "aktivny podiel")),
        Stlpec("scopus_id", "Scopus Author ID", to_str, False, ("scopus id", "scopus")),
        Stlpec("wos_id", "WoS ResearcherID", to_str, False, ("researcherid", "wos id", "researcher id")),
        Stlpec("orcid", "ORCID", to_str, False, ()),
    ], "Zoznam vysokoškolských učiteľov a ich prepočítaný úväzok."),
    "vyucba": TypImportu("vyucba", "Výučba", [
        *_ID_UCITELA,
        Stlpec("ak_rok", "Akademický rok", to_ak_rok, True, ("ak. rok", "akad. rok", "rok")),
        Stlpec("predmet", "Predmet", to_str, False, ("nazov predmetu", "kod predmetu")),
        Stlpec("jazyk", "Jazyk (SK/EN/MOB)", to_jazyk, False, ("jazyk", "jazyk vyucby")),
        Stlpec("odbor", "Študijný odbor", to_str, False, ("odbor", "studijny odbor")),
        Stlpec("hodiny", "Hodiny priamej výučby za ak. rok", to_float, True,
               ("hodiny", "pocet hodin", "odučené hodiny", "oducene hodiny", "rozsah")),
        Stlpec("pocet_studentov", "Počet študentov", to_float, False, ("studenti", "pocet zapisanych", "zapisani studenti")),
        Stlpec("studentohodiny", "Študentohodiny (nepovinné)", to_float, False, ("studentohodiny",)),
    ], "Priama výučba v akreditovaných študijných programoch (UIS)."),
    "zaverecne_prace": TypImportu("zaverecne_prace", "Záverečné práce", [
        *_ID_UCITELA,
        Stlpec("ak_rok", "Akademický rok", to_ak_rok, True, ("ak. rok", "rok obhajoby", "rok")),
        Stlpec("stupen", "Stupeň (Bc/Ing/PhD)", to_stupen, True, ("stupen", "typ prace", "druh prace", "typ zp")),
        Stlpec("pocet", "Počet", lambda v: to_float(v, 1.0), False, ()),
        Stlpec("nazov", "Názov práce", to_str, False, ("nazov", "tema")),
    ], "Úspešne obhájené záverečné práce – vedúci / školiteľ (UIS)."),
    "publikacie": TypImportu("publikacie", "Publikácie", [
        *_ID_UCITELA,
        Stlpec("rok", "Rok", to_int, True, ("rok vydania", "rok publikovania", "rok vykazania")),
        Stlpec("kategoria", "Kategória (tab. 3)", to_kategoria, True, ("kategoria", "skupina", "kategoria vystupu")),
        Stlpec("kvartil", "Kvartil", to_kvartil, False, ("kvartil", "q", "ais kvartil")),
        Stlpec("podiel", "Podiel (0–1)", to_podiel, False, ("podiel", "podiel autora", "podiel %")),
        Stlpec("nazov", "Názov", to_str, False, ("nazov", "titul")),
        Stlpec("zdroj", "Zdroj", to_str, False, ()),
        Stlpec("identifikator", "Identifikátor (DOI, EID, UT)", to_str, False, ("doi", "eid", "ut", "identifikator")),
    ], "Publikácie a podiely autorov (CREPČ / knižnica SPU)."),
    "projekty": TypImportu("projekty", "Projekty", [
        Stlpec("kod", "Kód projektu", to_str, True, ("kod", "cislo projektu", "evidencne cislo")),
        Stlpec("nazov", "Názov projektu", to_str, False, ("nazov",)),
        Stlpec("typ", "Typ projektu", to_str, False, ("typ", "program", "grantova schema")),
        Stlpec("rok", "Rok", to_int, True, ("kalendarny rok",)),
        Stlpec("suma", "Suma pripísaná SPU (€)", to_float, False, ("suma", "financie", "pridelene prostriedky", "eur")),
        Stlpec("kapacita_hodin", "Celková riešiteľská kapacita (h)", to_float, False, ("kapacita", "riesitelska kapacita")),
        Stlpec("pocet_riesitelov", "Počet riešiteľov (UIS)", to_int, False, ("pocet riesitelov",)),
        Stlpec("uis_id", "ID projektu v UIS", to_str, False, ("id projektu",)),
    ], "Projekty a finančné prostriedky pripísané na účet SPU (Sofia/SAP, UIS)."),
    "ucasti": TypImportu("ucasti", "Účasť na projektoch", [
        Stlpec("kod", "Kód projektu", to_str, True, ("kod", "cislo projektu")),
        Stlpec("rok", "Rok", to_int, True, ("kalendarny rok",)),
        *_ID_UCITELA,
        Stlpec("hodiny", "Hodiny za rok", to_float, True, ("hodiny", "vykazane hodiny", "pocet hodin")),
        Stlpec("zodpovedny", "Zodpovedný riešiteľ (áno/nie)", to_bool, False, ("zodpovedny", "rola", "zodpovedny riesitel")),
    ], "Vykázané hodiny učiteľov na projektoch (UIS)."),
}


def automaticke_mapovanie(typ: TypImportu, hlavicky: list[str]) -> dict[str, Optional[str]]:
    """Pre každé pole nájde stĺpec v súbore podľa názvu alebo synonyma."""
    nh = {norm(h): h for h in hlavicky if h is not None and str(h).strip()}
    mapovanie: dict[str, Optional[str]] = {}
    pouzite = set()
    for s in typ.stlpce:
        if s.pole in mapovanie:
            continue
        kandidati = [norm(s.hlavicka), norm(s.pole), *map(norm, s.synonyma)]
        found = None
        for k in kandidati:
            if k in nh and nh[k] not in pouzite:
                found = nh[k]
                break
        if found is None:
            for k in kandidati:
                for n, orig in nh.items():
                    if orig not in pouzite and len(k) >= 4 and (n.startswith(k) or k in n):
                        found = orig
                        break
                if found:
                    break
        mapovanie[s.pole] = found
        if found:
            pouzite.add(found)
    return mapovanie


# ------------------------------------------------------------------ čítanie súborov

def nacitaj_subor(path: str | Path, harok: Optional[str] = None) -> tuple[list[str], list[dict]]:
    """Načíta XLSX alebo CSV. Vráti (hlavičky, riadky ako dict)."""
    path = Path(path)
    if path.suffix.lower() in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1250", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:5000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
        rows = list(csv.reader(text.splitlines(), dialect))
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[harok] if harok and harok in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return [], []
    # hlavička = prvý riadok, v ktorom je aspoň polovica buniek vyplnená textom
    hi = 0
    for i, r in enumerate(rows[:15]):
        texty = [c for c in r if isinstance(c, str) and c.strip()]
        if len(texty) >= max(2, len([c for c in r if c not in (None, "")]) // 2):
            hi = i
            break
    hlavicky = [to_str(h) or f"Stĺpec {j + 1}" for j, h in enumerate(rows[hi])]
    data = []
    for r in rows[hi + 1:]:
        data.append({hlavicky[j]: (r[j] if j < len(r) else None) for j in range(len(hlavicky))})
    return hlavicky, data


def harky_suboru(path: str | Path) -> list[str]:
    path = Path(path)
    if path.suffix.lower() in (".csv", ".txt"):
        return []
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


# ------------------------------------------------------------------ import do databázy

@dataclass
class VysledokImportu:
    pridane: int = 0
    aktualizovane: int = 0
    preskocene: int = 0
    novi_ucitelia: list[str] = field(default_factory=list)
    chyby: list[str] = field(default_factory=list)

    def sprava(self) -> str:
        t = f"Pridané: {self.pridane}, aktualizované: {self.aktualizovane}, preskočené: {self.preskocene}."
        if self.novi_ucitelia:
            t += (f"\nAutomaticky založení učitelia ({len(self.novi_ucitelia)}) – doplňte im fakultu, ústav a úväzok: "
                  + ", ".join(self.novi_ucitelia[:15]) + (" …" if len(self.novi_ucitelia) > 15 else ""))
        if self.chyby:
            t += "\n\nUpozornenia:\n" + "\n".join(self.chyby[:20]) + ("\n…" if len(self.chyby) > 20 else "")
        return t


def _ucitel_pre_riadok(db: Databaza, rec: dict, res: VysledokImportu, zalozit: bool = True) -> Optional[int]:
    oc, meno = rec.get("osobne_cislo", ""), rec.get("meno", "")
    u = db.najdi_ucitela(oc, meno)
    if u:
        return u.id
    if not (oc or meno) or not zalozit:
        return None
    u = Ucitel(osobne_cislo=oc, meno=meno or oc)
    db.uloz(u, commit=False)
    res.novi_ucitelia.append(u.meno)
    return u.id


def importuj(db: Databaza, typ_kluc: str, riadky: list[dict], mapovanie: dict[str, Optional[str]],
             nahradit: bool = False, zdroj: str = "") -> VysledokImportu:
    """Importuje riadky podľa mapovania {pole: názov stĺpca v súbore}.

    nahradit=True: pred importom vymaže existujúce záznamy daného typu za roky obsiahnuté v súbore
    (zabraňuje duplicitám pri opakovanom importe toho istého exportu).
    """
    typ = TYPY[typ_kluc]
    res = VysledokImportu()
    zaznamy = []
    for i, row in enumerate(riadky, start=2):
        rec = {}
        for s in typ.stlpce:
            col = mapovanie.get(s.pole)
            rec[s.pole] = s.prevod(row.get(col) if col else None)
        chyba = [s.hlavicka for s in typ.stlpce if s.povinny and s.pole != "hodiny"
                 and (not mapovanie.get(s.pole) or rec.get(s.pole) in ("", None, 0))]
        if typ_kluc != "ucitelia" and typ_kluc != "projekty" and not (rec.get("osobne_cislo") or rec.get("meno")):
            chyba.append("učiteľ (osobné číslo alebo meno)")
        if chyba:
            res.preskocene += 1
            if any(v not in (None, "") for v in row.values()):
                res.chyby.append(f"Riadok {i}: chýba {', '.join(chyba)}")
            continue
        zaznamy.append(rec)

    if nahradit and typ_kluc in ("vyucba", "zaverecne_prace", "publikacie"):
        stlpec_roku = "ak_rok" if typ_kluc != "publikacie" else "rok"
        for r in {z[stlpec_roku] for z in zaznamy}:
            db.conn.execute(f"DELETE FROM {typ_kluc} WHERE {stlpec_roku} = ?", (r,))
    if nahradit and typ_kluc == "ucasti":
        for kod, rok in {(z["kod"], z["rok"]) for z in zaznamy}:
            p = db.najdi_projekt(kod, rok)
            if p:
                db.conn.execute("DELETE FROM ucasti WHERE projekt_id = ?", (p.id,))

    for rec in zaznamy:
        try:
            if typ_kluc == "ucitelia":
                u = db.najdi_ucitela(rec["osobne_cislo"], rec["meno"])
                novy = Ucitel(**{k: v for k, v in rec.items()})
                if u:
                    for s in typ.stlpce:   # aktualizuj iba namapované stĺpce
                        if mapovanie.get(s.pole):
                            setattr(u, s.pole, rec[s.pole])
                    db.uloz(u, commit=False)
                    res.aktualizovane += 1
                else:
                    db.uloz(novy, commit=False)
                    res.pridane += 1
            elif typ_kluc == "projekty":
                p = db.najdi_projekt(rec["kod"], rec["rok"])
                if p:
                    for s in typ.stlpce:
                        if mapovanie.get(s.pole):
                            setattr(p, s.pole, rec[s.pole])
                    db.uloz(p, commit=False)
                    res.aktualizovane += 1
                else:
                    db.uloz(Projekt(**rec), commit=False)
                    res.pridane += 1
            elif typ_kluc == "ucasti":
                p = db.najdi_projekt(rec["kod"], rec["rok"])
                if not p:
                    p = Projekt(kod=rec["kod"], rok=rec["rok"], nazov=rec["kod"])
                    db.uloz(p, commit=False)
                    res.chyby.append(f"Projekt {rec['kod']} ({rec['rok']}) nebol v zozname projektov – založený bez sumy.")
                uid = _ucitel_pre_riadok(db, rec, res)
                db.uloz(ProjektUcast(projekt_id=p.id, ucitel_id=uid, hodiny=rec["hodiny"],
                                     zodpovedny=rec["zodpovedny"]), commit=False)
                res.pridane += 1
            else:
                uid = _ucitel_pre_riadok(db, rec, res)
                cls = {"vyucba": Vyucba, "zaverecne_prace": ZaverecnaPraca, "publikacie": Publikacia}[typ_kluc]
                kw = {k: v for k, v in rec.items() if k not in ("osobne_cislo", "meno")}
                if typ_kluc == "publikacie" and not kw.get("zdroj"):
                    kw["zdroj"] = zdroj
                db.uloz(cls(ucitel_id=uid, **kw), commit=False)
                res.pridane += 1
        except Exception as e:  # noqa: BLE001 – chybný riadok nesmie zastaviť celý import
            res.preskocene += 1
            res.chyby.append(f"{rec}: {e}")
    db.commit()
    return res


# ------------------------------------------------------------------ uložené mapovania (profily pre UIS)

def _mapovania_path() -> Path:
    return config.app_data_dir() / "mapovania.json"


def nacitaj_profily() -> dict:
    p = _mapovania_path()
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            return {}
    return {}


def uloz_profil(nazov: str, typ_kluc: str, mapovanie: dict):
    prof = nacitaj_profily()
    prof[nazov] = {"typ": typ_kluc, "mapovanie": mapovanie}
    _mapovania_path().write_text(json.dumps(prof, ensure_ascii=False, indent=2), encoding="utf-8")


# ------------------------------------------------------------------ Excel šablóna / export údajov

_HLAVICKA_FILL = PatternFill("solid", fgColor="1E4620")
_HLAVICKA_FONT = Font(bold=True, color="FFFFFF")


def _validacia(ws, stlpec_idx: int, hodnoty_ref: str, n_riadkov: int = 2000):
    dv = DataValidation(type="list", formula1=hodnoty_ref, allow_blank=True)
    col = get_column_letter(stlpec_idx)
    dv.add(f"{col}2:{col}{n_riadkov}")
    ws.add_data_validation(dv)


def vytvor_sablonu(path: str | Path, db: Optional[Databaza] = None, params: Optional[dict] = None):
    """Vytvorí Excel šablónu. Ak je zadaná databáza, vyplní ju existujúcimi údajmi (záloha / zdieľanie)."""
    params = params or config.load_parametre()
    wb = Workbook()
    navod = wb.active
    navod.title = "Návod"
    riadky = [
        "Šablóna vstupných údajov – Hodnotenie pracovnej záťaže VŠ učiteľov SPU v Nitre",
        "Metodický pokyn 1/2023 v znení Dodatku č. 2 (účinný od 1. 7. 2025)",
        "",
        "Každý hárok zodpovedá jednému typu údajov. Hlavičky nemeňte, poradie riadkov je ľubovoľné.",
        "Učiteľa identifikuje Osobné číslo (ak chýba, Meno a priezvisko).",
        "Úväzok zadajte ako 0–1 (alebo v %), Podiel obdobia v PP = časť obdobia mimo materskej/rodičovskej (0–1).",
        "Jazyk výučby: SK, EN (aj SK/EN program vyučovaný v angličtine), MOB (mobilitní študenti).",
        "Akademický rok v tvare 2024/2025. Hodiny priamej výučby sa zadávajú za celý akademický rok.",
        "Podiel na publikácii zadajte ako 0–1 (alebo v %).",
        "Projekt sa zadáva za každý kalendárny rok samostatne (suma pripísaná SPU v danom roku).",
        "Celková riešiteľská kapacita: ak je prázdna, použije sa súčet hodín účastníkov.",
    ]
    for r in riadky:
        navod.append([r])
    navod["A1"].font = Font(bold=True, size=13)
    navod.column_dimensions["A"].width = 110

    cis = wb.create_sheet("Číselníky")
    zoznamy = {
        "Jazyk": list(config.JAZYKY_VYUCBY),
        "Stupeň": config.STUPNE_STUDIA,
        "Kategória": config.KATEGORIE_PUBLIKACII,
        "Kvartil": config.KVARTILY,
        "Odbor": list(params["koef_odbor"]),
        "Funkcia": config.FUNKCIE,
        "Typ projektu": config.TYPY_PROJEKTOV,
        "Áno/nie": ["áno", "nie"],
    }
    refs = {}
    for j, (k, vals) in enumerate(zoznamy.items(), start=1):
        cis.cell(row=1, column=j, value=k).font = Font(bold=True)
        for i, v in enumerate(vals, start=2):
            cis.cell(row=i, column=j, value=v)
        col = get_column_letter(j)
        refs[k] = f"'Číselníky'!${col}$2:${col}${len(vals) + 1}"
        cis.column_dimensions[col].width = max(14, max(len(str(v)) for v in vals) + 2)

    validacie = {"jazyk": "Jazyk", "stupen": "Stupeň", "kategoria": "Kategória", "kvartil": "Kvartil",
                 "odbor": "Odbor", "funkcia": "Funkcia", "typ": "Typ projektu", "zodpovedny": "Áno/nie"}

    data = db.nacitaj_vsetko() if db else None
    ucitelia = {u.id: u for u in data.ucitelia} if data else {}
    projekty = {p.id: p for p in data.projekty} if data else {}

    for kluc, typ in TYPY.items():
        ws = wb.create_sheet(typ.nazov)
        for j, s in enumerate(typ.stlpce, start=1):
            c = ws.cell(row=1, column=j, value=s.hlavicka)
            c.fill, c.font = _HLAVICKA_FILL, _HLAVICKA_FONT
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = max(14, len(s.hlavicka) + 2)
            if s.pole in validacie:
                _validacia(ws, j, refs[validacie[s.pole]])
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        if not data:
            continue
        for obj in getattr(data, kluc):
            row = []
            for s in typ.stlpce:
                if s.pole in ("osobne_cislo", "meno") and kluc not in ("ucitelia",):
                    u = ucitelia.get(getattr(obj, "ucitel_id", None))
                    row.append(getattr(u, s.pole, "") if u else "")
                elif kluc == "ucasti" and s.pole in ("kod", "rok"):
                    p = projekty.get(obj.projekt_id)
                    row.append(getattr(p, s.pole, "") if p else "")
                elif s.pole == "zodpovedny":
                    row.append("áno" if obj.zodpovedny else "nie")
                else:
                    row.append(getattr(obj, s.pole, ""))
            ws.append(row)
    wb.save(path)


def importuj_sablonu(db: Databaza, path: str | Path, nahradit: bool = True) -> dict[str, VysledokImportu]:
    """Importuje všetky hárky vyplnenej šablóny (v správnom poradí)."""
    wb_names = harky_suboru(path)
    out = {}
    for kluc in ("ucitelia", "projekty", "vyucba", "zaverecne_prace", "publikacie", "ucasti"):
        typ = TYPY[kluc]
        if typ.nazov not in wb_names:
            continue
        hlavicky, riadky = nacitaj_subor(path, typ.nazov)
        if not riadky:
            continue
        out[typ.nazov] = importuj(db, kluc, riadky, automaticke_mapovanie(typ, hlavicky), nahradit=nahradit, zdroj="Excel")
    return out
