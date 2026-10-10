"""Uloženie a načítanie všetkých údajov do pracovného priečinka so samostatnými zložkami.

    <priečinok>/
        01 Podmienky hodnotenia/   dokumenty metodiky, parametre.json, dokumenty.json
        02 Zoznam učiteľov/        ucitelia.xlsx
        03 Výučba/                 vyucba.xlsx
        04 Záverečné práce/        zaverecne_prace.xlsx
        05 Projekty/               projekty.xlsx (hárky Projekty a Účasti)
        06 Publikácie/             publikacie.xlsx
        07 Výsledky/               exporty výsledkov
        hodnotenie.json            verzia, dátum uloženia, sledované obdobie

Súbory sú bežné tabuľky Excelu – dajú sa otvoriť, skontrolovať aj upraviť. Hlavička stĺpca má tvar
„Popis [pole]“; podľa poľa v hranatých zátvorkách program súbor načíta späť bez straty údajov
(zachová aj väzby medzi učiteľmi, výučbou, projektmi …).
"""

from __future__ import annotations

import datetime as dt
import json
import re
import shutil
from dataclasses import fields
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .db import TABULKY, Databaza
from .version import __version__

ZLOZKY = {
    "podmienky": "01 Podmienky hodnotenia",
    "ucitelia": "02 Zoznam učiteľov",
    "vyucba": "03 Výučba",
    "zaverecne_prace": "04 Záverečné práce",
    "projekty": "05 Projekty",
    "publikacie": "06 Publikácie",
    "vysledky": "07 Výsledky",
}
SUBORY = {
    "ucitelia": ("ucitelia", "ucitelia.xlsx", "Učitelia"),
    "vyucba": ("vyucba", "vyucba.xlsx", "Výučba"),
    "zaverecne_prace": ("zaverecne_prace", "zaverecne_prace.xlsx", "Záverečné práce"),
    "projekty": ("projekty", "projekty.xlsx", "Projekty"),
    "ucasti": ("projekty", "projekty.xlsx", "Účasti na projektoch"),
    "publikacie": ("publikacie", "publikacie.xlsx", "Publikácie"),
}
POPISY = {
    "id": "ID", "ucitel_id": "ID učiteľa", "projekt_id": "ID projektu", "osobne_cislo": "Osobné číslo (UIS)",
    "titul_pred": "Titul pred", "meno": "Meno a priezvisko", "titul_za": "Titul za", "fakulta": "Fakulta",
    "ustav": "Ústav", "funkcia": "Funkcia", "uvazok": "Úväzok", "podiel_aktivny": "Podiel obdobia (staršie)",
    "scopus_id": "Scopus Author ID", "wos_id": "WoS ResearcherID", "orcid": "ORCID",
    "aktivny_vzdelavanie": "Aktívny – vzdelávanie", "aktivny_publikacie": "Aktívny – publikácie",
    "aktivny_projekty": "Aktívny – projekty", "poznamka": "Poznámka", "ak_rok": "Akademický rok",
    "predmet": "Predmet", "jazyk": "Jazyk", "odbor": "Študijný odbor", "hodiny": "Hodiny",
    "pocet_studentov": "Počet študentov", "studentohodiny": "Študentohodiny", "stupen": "Stupeň", "pocet": "Počet",
    "nazov": "Názov", "rok": "Rok", "kategoria": "Kategória", "kvartil": "Kvartil", "podiel": "Podiel",
    "zdroj": "Zdroj", "identifikator": "Identifikátor", "kod": "Kód", "typ": "Typ", "suma": "Suma SPU (€)",
    "kapacita_hodin": "Riešiteľská kapacita (h)", "pocet_riesitelov": "Počet riešiteľov", "uis_id": "ID v UIS",
    "zodpovedny": "Zodpovedný riešiteľ",
}
_HF = PatternFill("solid", fgColor="1E4620")


def zlozka(koren: Path, kluc: str) -> Path:
    p = Path(koren) / ZLOZKY[kluc]
    p.mkdir(parents=True, exist_ok=True)
    return p


def predvoleny_priecinok() -> Path:
    doc = Path.home() / "Documents"
    if not doc.exists():
        doc = Path.home()
    return doc / "SPU-Zataz údaje"


def je_priecinok_udajov(p: Path) -> bool:
    p = Path(p)
    return (p / "hodnotenie.json").exists() or (p / ZLOZKY["ucitelia"] / "ucitelia.xlsx").exists()


# ------------------------------------------------------------------ uloženie

def uloz(db: Databaza, koren: str | Path, nastavenia_hodnotenia: Optional[dict] = None,
         params: Optional[dict] = None) -> Path:
    koren = Path(koren)
    koren.mkdir(parents=True, exist_ok=True)
    for k in ZLOZKY:
        zlozka(koren, k)
    data = db.nacitaj_vsetko()
    mena = {u.id: u.cele_meno for u in data.ucitelia}
    projekty = {p.id: f"{p.kod} ({p.rok})" for p in data.projekty}

    zosity: dict[str, Workbook] = {}
    for tab, (zl, subor, harok) in SUBORY.items():
        wb = zosity.get(subor)
        if wb is None:
            wb = Workbook()
            wb.remove(wb.active)
            zosity[subor] = wb
        ws = wb.create_sheet(harok)
        polia = [f.name for f in fields(TABULKY[tab])]
        info = []
        if "ucitel_id" in polia:
            info.append(("Učiteľ (informatívne)", lambda o: mena.get(o.ucitel_id, "")))
        if "projekt_id" in polia:
            info.append(("Projekt (informatívne)", lambda o: projekty.get(o.projekt_id, "")))
        hlavicky = [f"{POPISY.get(p, p)} [{p}]" for p in polia] + [h for h, _ in info]
        ws.append(hlavicky)
        for j in range(1, len(hlavicky) + 1):
            c = ws.cell(row=1, column=j)
            c.fill, c.font = _HF, Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = 14 if j > 1 else 7
        for obj in getattr(data, tab):
            riadok = []
            for p in polia:
                v = getattr(obj, p)
                riadok.append(int(v) if isinstance(v, bool) else v)
            ws.append(riadok + [f(obj) for _, f in info])
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    for tab, (zl, subor, _) in SUBORY.items():
        if subor in zosity:
            zosity.pop(subor).save(zlozka(koren, zl) / subor)

    # podmienky hodnotenia
    pz = zlozka(koren, "podmienky")
    if params is not None:
        (pz / "parametre.json").write_text(json.dumps(params, ensure_ascii=False, indent=2), encoding="utf-8")
    from . import metodika
    src = metodika.priecinok()
    for f in src.iterdir():
        if f.is_file():
            try:
                shutil.copy2(f, pz / f.name)
            except (shutil.SameFileError, OSError):
                pass

    meta = {"program": "SPU-Zataz", "verzia": __version__, "ulozene": dt.datetime.now().isoformat(timespec="seconds"),
            "hodnotenie": nastavenia_hodnotenia or {},
            "pocty": {k: len(getattr(data, k)) for k in TABULKY}}
    (koren / "hodnotenie.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return koren


# ------------------------------------------------------------------ načítanie

def _prevod(hodnota, typ):
    t = str(typ)
    if hodnota is None or hodnota == "":
        return None
    if "bool" in t:
        return str(hodnota).strip().lower() in ("1", "true", "áno", "ano", "a", "x")
    if "int" in t and "Optional" not in t:
        try:
            return int(float(str(hodnota).replace(",", ".")))
        except ValueError:
            return None
    if "Optional[int]" in t:
        try:
            return int(float(hodnota))
        except (TypeError, ValueError):
            return None
    if "float" in t:
        try:
            return float(str(hodnota).replace(",", "."))
        except ValueError:
            return None
    return str(hodnota) if not isinstance(hodnota, str) else hodnota


def nacitaj(db: Databaza, koren: str | Path) -> dict:
    """Nahradí všetky údaje v databáze údajmi z priečinka. Vráti meta (hodnotenie.json) s počtami."""
    koren = Path(koren)
    if not (koren / ZLOZKY["ucitelia"] / "ucitelia.xlsx").exists():
        raise FileNotFoundError(f"V priečinku chýba {ZLOZKY['ucitelia']}/ucitelia.xlsx – nejde o uložené údaje programu.")
    objekty: dict[str, list] = {}
    for tab, (zl, subor, harok) in SUBORY.items():
        cesta = koren / ZLOZKY[zl] / subor
        objekty[tab] = []
        if not cesta.exists():
            continue
        wb = load_workbook(cesta, read_only=True, data_only=True)
        if harok not in wb.sheetnames:
            continue
        ws = wb[harok]
        riadky = ws.iter_rows(values_only=True)
        hlavicka = next(riadky, None) or ()
        mapa = {}
        for j, h in enumerate(hlavicka):
            m = re.search(r"\[(\w+)\]\s*$", str(h or ""))
            if m:
                mapa[j] = m.group(1)
        cls = TABULKY[tab]
        typy = {f.name: f.type for f in fields(cls)}
        for r in riadky:
            if not r or all(v in (None, "") for v in r):
                continue
            kw = {}
            for j, pole in mapa.items():
                if pole in typy and j < len(r):
                    v = _prevod(r[j], typy[pole])
                    if v is not None:
                        kw[pole] = v
            if "id" not in kw:
                continue
            objekty[tab].append(cls(**kw))
        wb.close()
    db.vymaz_vsetko()
    for tab in ("ucitelia", "projekty", "vyucba", "zaverecne_prace", "publikacie", "ucasti"):
        for o in objekty[tab]:
            db.vloz_s_id(o)
    db.commit()
    meta = {}
    mp = koren / "hodnotenie.json"
    if mp.exists():
        try:
            meta = json.loads(mp.read_text(encoding="utf-8"))
        except ValueError:
            meta = {}
    meta["nacitane"] = {k: len(v) for k, v in objekty.items()}
    return meta


def nacitaj_podmienky(koren: str | Path) -> Optional[dict]:
    """Skopíruje dokumenty metodiky z priečinka údajov do programu a vráti uložené parametre (alebo None)."""
    from . import metodika
    pz = Path(koren) / ZLOZKY["podmienky"]
    if not pz.exists():
        return None
    for f in pz.iterdir():
        if f.is_file() and f.name != "parametre.json":
            try:
                shutil.copy2(f, metodika.priecinok() / f.name)
            except (shutil.SameFileError, OSError):
                pass
    p = pz / "parametre.json"
    if p.exists():
        try:
            return config._merge(config.DEFAULT_PARAMETRE, json.loads(p.read_text(encoding="utf-8")))
        except ValueError:
            return None
    return None
